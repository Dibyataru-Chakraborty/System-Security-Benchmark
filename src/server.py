import os
import sys
import re
import csv
import json
import tempfile
import subprocess
import platform
import winreg as wrg
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file

# Handle PyInstaller frozen executable vs source paths
if getattr(sys, "frozen", False):
    BUNDLE_DIR = getattr(sys, "_MEIPASS", os.path.dirname(sys.executable))
    PROJECT_ROOT = os.path.dirname(sys.executable)
else:
    BUNDLE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    PROJECT_ROOT = BUNDLE_DIR

if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)
if BUNDLE_DIR not in sys.path:
    sys.path.insert(0, BUNDLE_DIR)

import src.benchmark_parser as bp
from src.excel_exporter import export_audit_results_to_excel, export_benchmark_rules_to_excel
import src.dashboard_handel as dh

app = Flask(
    __name__,
    template_folder=os.path.join(BUNDLE_DIR, "templates"),
    static_folder=os.path.join(BUNDLE_DIR, "static")
)

# Global in-memory cache
STATE = {
    "rules": [],
    "audit_records": [],
    "pdf_path": os.path.join(PROJECT_ROOT, "CIS_Microsoft_Windows_11_Enterprise_Benchmark_v5.1.0.pdf"),
    "session_cookie": None,
    "compliance_target": "CIS Security Benchmark",
    "evaluation_engine": "Risk Nexus CIS Security Benchmark Engine",
    "benchmark_title": "",
    "benchmark_version": "",
    "active_hive": "HKEY_LOCAL_MACHINE",
    "audit_timestamp": ""
}

OUTPUTS_DIR = os.path.join(PROJECT_ROOT, "outputs")
os.makedirs(OUTPUTS_DIR, exist_ok=True)

# Load saved cookie if available
cookie_file = os.path.join(OUTPUTS_DIR, "session_cookie.txt")
if os.path.exists(cookie_file):
    try:
        with open(cookie_file, "r", encoding="utf-8") as f:
            STATE["session_cookie"] = f.read().strip()
    except Exception:
        pass

# -------------------------------------------------------------
# Web Page Route
# -------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")

# -------------------------------------------------------------
# System Info API (Ultra-fast native Windows telemetry)
# -------------------------------------------------------------
@app.route("/api/system-info", methods=["GET"])
def get_system_info():
    system_info = [
        {"key": "Host Name", "value": platform.node() or "Local Machine"},
        {"key": "OS Name", "value": f"{platform.system()} {platform.release()}"},
        {"key": "OS Version / Kernel", "value": platform.version()},
        {"key": "Architecture", "value": platform.machine()},
        {"key": "Processor", "value": os.environ.get("PROCESSOR_IDENTIFIER", platform.processor() or "x86_64")},
        {"key": "Processor Cores", "value": os.environ.get("NUMBER_OF_PROCESSORS", "1")}
    ]

    try:
        with wrg.OpenKey(wrg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion") as k:
            for val_name, label in [
                ("ProductName", "Product Name"),
                ("DisplayVersion", "Display Version"),
                ("CurrentBuild", "OS Build"),
                ("EditionID", "Windows Edition"),
                ("RegisteredOwner", "Registered Owner")
            ]:
                try:
                    val, _ = wrg.QueryValueEx(k, val_name)
                    if val:
                        system_info.append({"key": label, "value": str(val)})
                except FileNotFoundError:
                    pass
    except Exception:
        pass

    return jsonify({
        "success": True,
        "data": system_info,
        "hostname": platform.node(),
        "os_name": f"{platform.system()} {platform.release()} ({platform.version()})",
        "arch": platform.machine()
    })

# -------------------------------------------------------------
# Technologies & Online Benchmarks API (Dynamic from CIS)
# -------------------------------------------------------------
@app.route("/api/technologies", methods=["GET"])
def get_technologies():
    try:
        import Cis_Download
        raw_techs = Cis_Download.get_technologies()
        flat_techs = Cis_Download.flatten_technologies(raw_techs)
        
        formatted = []
        if flat_techs:
            for t in flat_techs:
                formatted.append({
                    "id": str(t.get("id", "")),
                    "name": t.get("title", "Unknown"),
                    "category": t.get("category", "General"),
                    "version": t.get("latest_technology_version", "")
                })
            formatted.sort(key=lambda x: x["name"])
        return jsonify({"success": True, "technologies": formatted})
    except Exception as e:
        print(f"Error fetching technologies: {e}")
        return jsonify({"success": False, "technologies": [], "error": str(e)}), 500

@app.route("/api/benchmarks/<tech_id>", methods=["GET"])
def get_benchmarks(tech_id):
    try:
        import Cis_Download
        bms = Cis_Download.get_latest_benchmarks(int(tech_id))
        out = []
        if bms and isinstance(bms, list):
            for b in bms:
                b_title = b.get("title", "CIS Benchmark")
                b_ver = b.get("version", "1.0.0")
                docs = b.get("documents", [])
                for doc in docs:
                    doc_title = doc.get("title", b_title)
                    pid = doc.get("pardot-id", "")
                    fname = doc.get("filename", f"{b_title}.pdf")
                    if pid:
                        clean_pid = pid.lstrip("/")
                        download_url = f"https://learn.cisecurity.org/{clean_pid}"
                        out.append({
                            "title": f"{doc_title} (v{b_ver})",
                            "pardot_id": pid,
                            "download_url": download_url,
                            "filename": fname
                        })
        return jsonify({"success": True, "benchmarks": out})
    except Exception as e:
        print(f"Error fetching benchmarks for tech {tech_id}: {e}")
        return jsonify({"success": False, "benchmarks": [], "error": str(e)}), 500

# -------------------------------------------------------------
# Download Benchmark PDF Online API
# -------------------------------------------------------------
@app.route("/api/download-benchmark", methods=["POST"])
def download_benchmark():
    data = request.get_json() or {}
    pardot_id = data.get("pardot_id", "c/cis-microsoft-windows-11-enterprise-benchmark")
    filename = data.get("filename", "CIS_Microsoft_Windows_11_Enterprise_Benchmark_v5.1.0.pdf")
    cookie = data.get("session_cookie") or STATE.get("session_cookie")

    dest_dir = tempfile.gettempdir()
    dest = os.path.join(dest_dir, filename)

    if pardot_id.startswith("http://") or pardot_id.startswith("https://"):
        download_url = pardot_id
    else:
        clean_id = pardot_id.lstrip("/")
        download_url = f"https://learn.cisecurity.org/{clean_id}"

    try:
        try:
            from curl_cffi import requests as http_lib
            use_curl = True
        except ImportError:
            import requests as http_lib
            use_curl = False

        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/139.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
            "Referer": "https://www.cisecurity.org/",
            "Sec-Ch-Ua": '"Not A(Brand";v="8", "Chromium";v="132", "Google Chrome";v="132"',
            "Sec-Ch-Ua-Mobile": "?0",
            "Sec-Ch-Ua-Platform": '"Windows"',
        }

        session = http_lib.Session(impersonate="chrome") if use_curl else http_lib.Session()
        session.headers.update(headers)
        if cookie:
            session.cookies.set("laravel_session", cookie, domain=".cisecurity.org")

        resp = session.get(download_url, timeout=30, allow_redirects=True)
        
        # Check if direct PDF content
        if resp.status_code == 200 and (resp.content.startswith(b"%PDF-") or "pdf" in resp.headers.get("Content-Type", "").lower()):
            with open(dest, "wb") as f:
                f.write(resp.content)
            STATE["pdf_path"] = dest
            return jsonify({
                "success": True,
                "path": dest,
                "download_url": download_url,
                "message": "Benchmark downloaded successfully!"
            })
        
        # If HTML returned, look for direct PDF link inside
        if resp.status_code == 200 and b"<html" in resp.content.lower():
            import re
            pdf_links = re.findall(r'href=[\'"](https?://[^\'"]+\.pdf[^\'"]*)[\'"]', resp.text, re.IGNORECASE)
            if pdf_links:
                target_pdf_url = pdf_links[0]
                pdf_resp = session.get(target_pdf_url, timeout=30, allow_redirects=True)
                if pdf_resp.status_code == 200 and pdf_resp.content.startswith(b"%PDF-"):
                    with open(dest, "wb") as f:
                        f.write(pdf_resp.content)
                    STATE["pdf_path"] = dest
                    return jsonify({
                        "success": True,
                        "path": dest,
                        "download_url": target_pdf_url,
                        "message": "Benchmark downloaded via direct link!"
                    })
    except Exception as e:
        print(f"Online download error: {e}")

    # Fallback to local copy if available
    local_copy = os.path.join(PROJECT_ROOT, "CIS_Microsoft_Windows_11_Enterprise_Benchmark_v5.1.0.pdf")
    if os.path.exists(local_copy):
        import shutil
        shutil.copy2(local_copy, dest)
        STATE["pdf_path"] = dest
        return jsonify({
            "success": True,
            "path": dest,
            "download_url": download_url,
            "message": "Using bundled CIS Benchmark copy (Online download gated by CIS authentication)."
        })

    return jsonify({
        "success": False,
        "download_url": download_url,
        "error": "Failed to auto-download. Please use the 'Open in Browser' button to download directly."
    }), 400

# -------------------------------------------------------------
# Upload PDF Benchmark API
# -------------------------------------------------------------
@app.route("/api/upload-pdf", methods=["POST"])
def upload_pdf():
    if "file" not in request.files:
        return jsonify({"success": False, "error": "No file uploaded"}), 400
    file = request.files["file"]
    if file.filename == "":
        return jsonify({"success": False, "error": "No file selected"}), 400

    upload_dir = os.path.join(PROJECT_ROOT, "uploads")
    os.makedirs(upload_dir, exist_ok=True)
    save_path = os.path.join(upload_dir, file.filename)
    file.save(save_path)
    STATE["pdf_path"] = save_path

    # Extract metadata from filename/pdf
    meta = bp.extract_benchmark_metadata(pdf_path=save_path)
    STATE["compliance_target"] = meta.get("compliance_target")
    STATE["evaluation_engine"] = meta.get("evaluation_engine")
    STATE["benchmark_title"] = meta.get("title")
    STATE["benchmark_version"] = meta.get("version")

    return jsonify({
        "success": True,
        "path": save_path,
        "compliance_target": meta.get("compliance_target"),
        "evaluation_engine": meta.get("evaluation_engine")
    })

@app.route("/api/parse-benchmark", methods=["POST"])
def parse_benchmark():
    data = request.get_json() or {}
    pdf_path = data.get("pdf_path") or STATE.get("pdf_path")
    rules = []

    if pdf_path and os.path.exists(pdf_path):
        try:
            rules = bp.parse_cis_pdf(pdf_path)
            STATE["rules"] = rules
            STATE["pdf_path"] = pdf_path
            
            # Extract dynamic metadata
            meta = bp.extract_benchmark_metadata(pdf_path=pdf_path)
            if data.get("benchmark_title"):
                meta["compliance_target"] = data["benchmark_title"]
                meta["title"] = data["benchmark_title"]
                meta["evaluation_engine"] = f"Digi Samurai CIS Security Benchmark Engine ({meta.get('version', 'v1.0.0')})"
            
            STATE["compliance_target"] = meta.get("compliance_target")
            STATE["evaluation_engine"] = meta.get("evaluation_engine")
            STATE["benchmark_title"] = meta.get("title")
            STATE["benchmark_version"] = meta.get("version")
        except Exception as e:
            return jsonify({"success": False, "error": f"Error parsing PDF: {e}"}), 500
    else:
        # Fallback to pre-cached benchmark rules Excel if available
        rules_excel = os.path.join(OUTPUTS_DIR, "CIS_Windows_11_Hardening_Rules.xlsx")
        if os.path.exists(rules_excel):
            import openpyxl
            wb = openpyxl.load_workbook(rules_excel)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[1]:
                    rules.append({
                        "Rule ID": row[1],
                        "Profile/Level": row[2] if len(row) > 2 else "L1",
                        "Rule Title": row[3] if len(row) > 3 else "",
                        "Hive": row[4] if len(row) > 4 else "HKLM",
                        "Key Path": row[5] if len(row) > 5 else "",
                        "Value Name": row[6] if len(row) > 6 else "",
                        "Value Type": row[7] if len(row) > 7 else "REG_DWORD",
                        "Recommended Value": str(row[8]) if len(row) > 8 else "1",
                        "Full Registry Path": row[9] if len(row) > 9 else f"{row[4]}\\{row[5]}:{row[6]}"
                    })
            STATE["rules"] = rules
            if data.get("benchmark_title"):
                STATE["compliance_target"] = data["benchmark_title"]
                STATE["evaluation_engine"] = "Digi Samurai CIS Security Benchmark Engine"
        else:
            return jsonify({
                "success": False,
                "error": "No benchmark PDF or cached rules found. Please download or browse for a CIS Benchmark PDF."
            }), 404

    # Breakdown by Hive & Level
    hives = {}
    levels = {}
    for r in rules:
        h = r.get("Hive", "HKLM")
        lvl = r.get("Profile/Level", "L1")
        hives[h] = hives.get(h, 0) + 1
        levels[lvl] = levels.get(lvl, 0) + 1

    return jsonify({
        "success": True,
        "total_rules": len(rules),
        "hives": hives,
        "levels": levels,
        "pdf_path": pdf_path if (pdf_path and os.path.exists(pdf_path)) else "Active Benchmark Rules",
        "compliance_target": STATE.get("compliance_target"),
        "evaluation_engine": STATE.get("evaluation_engine"),
        "rules": rules[:50]
    })

# -------------------------------------------------------------
# Live Security Audit API
# -------------------------------------------------------------
def format_and_compare_registry_value(val_data, rec_val, val_type="REG_DWORD"):
    """
    Normalizes Windows registry data (handling lists, REG_MULTI_SZ, binary, and integers)
    and evaluates compliance against the CIS recommendation cleanly.
    """
    raw_elements = []
    if isinstance(val_data, (list, tuple)):
        cleaned_list = [str(x).strip() for x in val_data if str(x).strip()]
        if not cleaned_list:
            display_val = "None"
        elif len(cleaned_list) == 1:
            display_val = cleaned_list[0]
        else:
            display_val = ", ".join(cleaned_list)
        raw_elements = [str(x).strip().lower() for x in val_data if str(x).strip()]
    elif isinstance(val_data, bytes):
        try:
            if len(val_data) <= 8:
                int_val = int.from_bytes(val_data, byteorder="little")
                display_val = str(int_val)
            else:
                display_val = val_data.hex()
        except Exception:
            display_val = val_data.hex()
    elif val_data is None:
        display_val = "Not Configured"
    else:
        display_val = str(val_data).strip()

    rec_str = str(rec_val).strip() if rec_val is not None else ""
    rec_lower = rec_str.lower()
    
    # Strip any accidental brackets or quotes
    rec_unwrapped = re.sub(r"^[\(\[\'\"\s]+|[\)\]\'\"\s]+$", "", rec_lower)
    curr_unwrapped = re.sub(r"^[\(\[\'\"\s]+|[\)\]\'\"\s]+$", "", display_val.lower())

    is_compliant = False

    if rec_lower == display_val.lower() or (rec_unwrapped and rec_unwrapped == curr_unwrapped):
        is_compliant = True
    elif ("\\" in rec_unwrapped or "/" in rec_unwrapped) and ("\\" in curr_unwrapped or "/" in curr_unwrapped):
        try:
            norm_rec = os.path.normpath(os.path.expandvars(rec_unwrapped)).lower()
            norm_curr = os.path.normpath(os.path.expandvars(curr_unwrapped)).lower()
            if norm_rec == norm_curr:
                is_compliant = True
        except Exception:
            pass
    elif rec_unwrapped in ("1", "enabled", "enable", "true") and curr_unwrapped in ("1", "enabled", "enable", "true"):
        is_compliant = True
    elif rec_unwrapped in ("0", "disabled", "disable", "false") and curr_unwrapped in ("0", "disabled", "disable", "false"):
        is_compliant = True
    elif rec_unwrapped in ("", "none", "null", "not configured") and curr_unwrapped in ("", "none", "null", "not configured", "[]"):
        is_compliant = True
    elif raw_elements:
        if rec_unwrapped in raw_elements:
            is_compliant = True
        elif rec_unwrapped in ("none", "") and len(raw_elements) == 0:
            is_compliant = True
        else:
            rec_items = [x.strip().lower() for x in re.split(r"[,;]+", rec_unwrapped) if x.strip()]
            if rec_items and all(item in raw_elements for item in rec_items):
                is_compliant = True
    else:
        try:
            rec_num = int(rec_unwrapped, 0)
            curr_num = int(curr_unwrapped, 0)
            if rec_num == curr_num:
                is_compliant = True
        except (ValueError, TypeError):
            pass

    return display_val, is_compliant

@app.route("/api/audit", methods=["POST", "GET"])
def run_audit(selected_hive=None):
    try:
        data = {}
        if selected_hive is None:
            try:
                data = request.get_json(silent=True) or {}
                selected_hive = data.get("hive", "HKEY_LOCAL_MACHINE")
            except Exception:
                selected_hive = "HKEY_LOCAL_MACHINE"

        STATE["active_hive"] = selected_hive
        STATE["audit_timestamp"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        if data.get("benchmark_title"):
            STATE["compliance_target"] = data["benchmark_title"]
            STATE["benchmark_title"] = data["benchmark_title"]
            STATE["evaluation_engine"] = f"Digi Samurai CIS Security Benchmark Engine ({STATE.get('benchmark_version', 'v1.0.0')})"

        hive_map = {
            "HKEY_LOCAL_MACHINE": "HKLM",
            "HKEY_CURRENT_USER": "HKCU",
            "HKEY_CLASSES_ROOT": "HKCR",
            "HKEY_USERS": "HKU",
            "HKEY_CURRENT_CONFIG": "HKCC"
        }
        hive_abbr = hive_map.get(selected_hive, "HKLM")

        # If rules not parsed yet, discover and parse PDF or load from existing benchmark rules Excel
        if not STATE.get("rules"):
            candidate_pdf = data.get("pdf_path") or STATE.get("pdf_path")
            
            # If candidate_pdf not set or doesn't exist, search common directories
            if not candidate_pdf or not os.path.exists(candidate_pdf):
                search_dirs = [PROJECT_ROOT, OUTPUTS_DIR, os.path.join(PROJECT_ROOT, "uploads"), os.getcwd()]
                for sdir in search_dirs:
                    if os.path.exists(sdir):
                        for fname in os.listdir(sdir):
                            if fname.lower().endswith(".pdf"):
                                full_p = os.path.join(sdir, fname)
                                if os.path.isfile(full_p):
                                    candidate_pdf = full_p
                                    break
                    if candidate_pdf and os.path.exists(candidate_pdf):
                        break

            if candidate_pdf and os.path.exists(candidate_pdf):
                print(f"[*] Auto-parsing benchmark PDF: {candidate_pdf}")
                STATE["rules"] = bp.parse_cis_pdf(candidate_pdf)
                STATE["pdf_path"] = candidate_pdf
                meta = bp.extract_benchmark_metadata(pdf_path=candidate_pdf)
                if not data.get("benchmark_title"):
                    STATE["compliance_target"] = meta.get("compliance_target")
                    STATE["evaluation_engine"] = meta.get("evaluation_engine")
                    STATE["benchmark_title"] = meta.get("title")
                    STATE["benchmark_version"] = meta.get("version")
            else:
                rules_excel = os.path.join(OUTPUTS_DIR, "CIS_Windows_11_Hardening_Rules.xlsx")
                if os.path.exists(rules_excel):
                    import openpyxl
                    wb = openpyxl.load_workbook(rules_excel)
                    ws = wb.active
                    loaded_rules = []
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if row and len(row) > 1 and row[1]:
                            loaded_rules.append({
                                "Rule ID": str(row[1]),
                                "Profile/Level": str(row[2]) if len(row) > 2 else "L1",
                                "Severity": str(row[3]) if len(row) > 3 else "Medium",
                                "Rule Title": str(row[4]) if len(row) > 4 else "",
                                "Hive": str(row[5]) if len(row) > 5 else "HKLM",
                                "Key Path": str(row[6]) if len(row) > 6 else "",
                                "Value Name": str(row[7]) if len(row) > 7 else "",
                                "Value Type": str(row[8]) if len(row) > 8 else "REG_DWORD",
                                "Recommended Value": str(row[9]) if len(row) > 9 else "1",
                                "Full Registry Path": str(row[10]) if len(row) > 10 else f"{row[5]}\\{row[6]}:{row[7]}"
                            })
                    STATE["rules"] = loaded_rules

        rules_to_audit = [r for r in STATE.get("rules", []) if r.get("Hive") == hive_abbr]
        if not rules_to_audit and STATE.get("rules"):
            rules_to_audit = STATE["rules"]

        if not rules_to_audit:
            return jsonify({
                "success": False,
                "error": "No benchmark rules found. Please select or upload a CIS Benchmark PDF in Step 1 first."
            }), 400

        location_map = {
            "HKLM": wrg.HKEY_LOCAL_MACHINE,
            "HKCU": wrg.HKEY_CURRENT_USER,
            "HKCR": wrg.HKEY_CLASSES_ROOT,
            "HKU": wrg.HKEY_USERS,
            "HKCC": wrg.HKEY_CURRENT_CONFIG
        }

        output_records = []
        pass_count = 0

        for idx, r in enumerate(rules_to_audit, 1):
            reg_path = r.get("Full Registry Path", "")
            rule_id = r.get("Rule ID", f"CIS-{idx}")
            level = r.get("Profile/Level", "L1")
            title = r.get("Rule Title", "Registry Rule")
            val_type = r.get("Value Type", "REG_DWORD")
            rec_val = str(r.get("Recommended Value", "1"))
            hive = r.get("Hive", "HKLM")
            locate = r.get("Key Path", "")
            value = r.get("Value Name", "")
            severity = r.get("Severity") or r.get("severity") or bp.determine_severity(rule_id, level, title, locate)

            locate_upper = locate.upper()
            category = "s/w" if "SOFTWARE" in locate_upper else "sys" if "SYSTEM" in locate_upper else hive
            if "POLICIES" in locate_upper: subcategory = "policies"
            elif "CLASSES" in locate_upper: subcategory = "classes"
            elif "MICROSOFT" in locate_upper: subcategory = "microsoft"
            elif "CONTROL" in locate_upper: subcategory = "control"
            elif "SERVICES" in locate_upper: subcategory = "services"
            else: subcategory = "other"

            val_exists = False
            raw_val_data = None
            current_val = "Not Configured (Missing)"
            try:
                loc_key = location_map.get(hive, wrg.HKEY_LOCAL_MACHINE)
                with wrg.OpenKey(loc_key, locate, 0, wrg.KEY_READ) as search_key:
                    try:
                        raw_val_data, _ = wrg.QueryValueEx(search_key, value)
                        val_exists = True
                    except FileNotFoundError:
                        current_val = "Key exists, Value missing"
            except FileNotFoundError:
                current_val = "Not Configured (Missing)"
            except PermissionError:
                current_val = "Access Denied"
            except Exception as e:
                current_val = f"Error: {e}"

            found = "No"
            status = "FAIL"
            if val_exists:
                formatted_curr, is_compliant = format_and_compare_registry_value(raw_val_data, rec_val, val_type)
                current_val = formatted_curr
                if is_compliant:
                    status = "PASS"
                    found = "Yes"
                    pass_count += 1
                else:
                    status = "FAIL"
                    found = "No"

            output_records.append({
                "sl": idx,
                "rule_id": rule_id,
                "level": level,
                "severity": severity,
                "title": title,
                "hive": hive,
                "key_path": locate,
                "value_name": value,
                "value_type": val_type,
                "recommended_value": rec_val,
                "current_value": current_val,
                "status": status,
                "found": found,
                "category": category,
                "subcategory": subcategory,
                "path": reg_path
            })

        STATE["audit_records"] = output_records

        # 1. Save CSV (The only automated file created on audit)
        csv_file = os.path.join(OUTPUTS_DIR, "registry_check_results.csv")
        csv_header = ["sl", "path", "found", "rule_id", "level", "severity", "title", "current_value", "recommended_value", "status", "category", "subcategory"]
        csv_rows = []
        for r in output_records:
            csv_rows.append([
                r["sl"], r["path"], r["found"], r["rule_id"], r["level"], r["severity"], r["title"],
                r["current_value"], r["recommended_value"], r["status"], r["category"], r["subcategory"]
            ])
        with open(csv_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(csv_header)
            writer.writerows(csv_rows)

        total_rules = len(output_records)
        fail_count = total_rules - pass_count
        score_pct = round((pass_count / total_rules * 100.0), 1) if total_rules > 0 else 0.0

        return jsonify({
            "success": True,
            "total_rules": total_rules,
            "passed": pass_count,
            "failed": fail_count,
            "score_pct": score_pct,
            "records": output_records
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": f"Live audit error: {str(e)}"}), 500

# -------------------------------------------------------------
# System Hardening API
# -------------------------------------------------------------
@app.route("/api/harden", methods=["POST"])
def apply_hardening():
    data = request.get_json() or {}
    paths_to_harden = data.get("paths", [])
    do_backup = data.get("backup", True)

    if not paths_to_harden:
        return jsonify({"success": False, "error": "No registry paths selected for hardening."}), 400

    location_map = {
        "HKLM": wrg.HKEY_LOCAL_MACHINE,
        "HKEY_LOCAL_MACHINE": wrg.HKEY_LOCAL_MACHINE,
        "HKCU": wrg.HKEY_CURRENT_USER,
        "HKEY_CURRENT_USER": wrg.HKEY_CURRENT_USER,
        "HKCR": wrg.HKEY_CLASSES_ROOT,
        "HKEY_CLASSES_ROOT": wrg.HKEY_CLASSES_ROOT,
        "HKU": wrg.HKEY_USERS,
        "HKEY_USERS": wrg.HKEY_USERS,
        "HKCC": wrg.HKEY_CURRENT_CONFIG,
        "HKEY_CURRENT_CONFIG": wrg.HKEY_CURRENT_CONFIG
    }

    # Backup if requested
    backup_path = None
    if do_backup:
        backup_path = os.path.join(OUTPUTS_DIR, "backup_registry.reg")
        create_backup_file(paths_to_harden, backup_path)

    # Build rules lookup
    all_rules = STATE.get("rules") or []
    if not all_rules:
        rules_excel = os.path.join(OUTPUTS_DIR, "CIS_Windows_11_Hardening_Rules.xlsx")
        if os.path.exists(rules_excel):
            import openpyxl
            wb = openpyxl.load_workbook(rules_excel)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[1]:
                    all_rules.append({
                        "Rule ID": row[1],
                        "Profile/Level": row[2] if len(row) > 2 else "L1",
                        "Rule Title": row[3] if len(row) > 3 else "",
                        "Hive": row[4] if len(row) > 4 else "HKLM",
                        "Key Path": row[5] if len(row) > 5 else "",
                        "Value Name": row[6] if len(row) > 6 else "",
                        "Value Type": row[7] if len(row) > 7 else "REG_DWORD",
                        "Recommended Value": str(row[8]) if len(row) > 8 else "1",
                        "Full Registry Path": row[9] if len(row) > 9 else f"{row[4]}\\{row[5]}:{row[6]}"
                    })

    rules_by_path = {}
    for r in all_rules:
        fp = r.get("Full Registry Path", "")
        if fp:
            rules_by_path[fp.strip().lower()] = r
            rules_by_path[fp.replace("\\", "").strip().lower()] = r

    success_count = 0
    fail_count = 0
    denied_count = 0
    log_entries = []

    for raw_path in paths_to_harden:
        reg_path = str(raw_path).replace("/", "\\").strip()
        try:
            if ":" not in reg_path:
                fail_count += 1
                log_entries.append({"path": raw_path, "status": "ERROR", "message": "Invalid format: missing ':' between key path and value name"})
                continue

            full_key, value = reg_path.rsplit(":", 1)
            
            # Extract hive and subkey path
            if "\\" in full_key:
                location, locate = full_key.split("\\", 1)
            else:
                # Handle edge case where backslashes were stripped
                for h in ["HKLM", "HKEY_LOCAL_MACHINE", "HKCU", "HKEY_CURRENT_USER", "HKCR", "HKU", "HKCC"]:
                    if full_key.upper().startswith(h):
                        location = h
                        locate = full_key[len(h):]
                        break
                else:
                    location = "HKLM"
                    locate = full_key

            loc_upper = location.upper().strip()
            if loc_upper in location_map:
                loc_key = location_map[loc_upper]
                rule = rules_by_path.get(reg_path.lower()) or rules_by_path.get(reg_path.replace("\\", "").lower()) or {}
                val_type = str(rule.get("Value Type", "REG_DWORD")).upper()
                rec_val = str(rule.get("Recommended Value", "1")).strip()

                # Determine WinReg Type & Conversion
                if "SZ" in val_type and "MULTI" not in val_type and "EXPAND" not in val_type:
                    reg_type_enum = wrg.REG_SZ
                    val_to_set = rec_val
                elif "MULTI" in val_type:
                    reg_type_enum = wrg.REG_MULTI_SZ
                    clean_rec = re.sub(r"^[\(\[\'\"\s]+|[\)\]\'\"\s]+$", "", rec_val)
                    if clean_rec.lower() in ("none", "", "null", "disabled", "0"):
                        val_to_set = []
                    else:
                        val_to_set = [s.strip() for s in re.split(r"[,;]+", clean_rec) if s.strip()]
                elif "EXPAND" in val_type:
                    reg_type_enum = wrg.REG_EXPAND_SZ
                    val_to_set = rec_val
                elif "QWORD" in val_type:
                    reg_type_enum = wrg.REG_QWORD
                    try:
                        val_to_set = int(rec_val, 0) if rec_val.startswith("0x") else int(rec_val)
                    except ValueError:
                        val_to_set = 1
                else:  # Default REG_DWORD
                    reg_type_enum = wrg.REG_DWORD
                    try:
                        val_to_set = int(rec_val, 0) if rec_val.startswith("0x") else int(rec_val)
                    except ValueError:
                        val_to_set = 0 if rec_val in ("0", "disabled", "disable", "false") else 1

                try:
                    with wrg.CreateKey(loc_key, locate) as key:
                        wrg.SetValueEx(key, value, 0, reg_type_enum, val_to_set)
                    success_count += 1
                    log_entries.append({
                        "path": raw_path,
                        "status": "SUCCESS",
                        "message": f"Applied value: {val_to_set} ({val_type})"
                    })
                except PermissionError:
                    denied_count += 1
                    log_entries.append({"path": raw_path, "status": "DENIED", "message": "Permission denied (Run as Administrator)"})
                except Exception as e:
                    fail_count += 1
                    log_entries.append({"path": raw_path, "status": "FAILED", "message": str(e)})
            else:
                fail_count += 1
                log_entries.append({"path": raw_path, "status": "ERROR", "message": f"Unknown hive: {location}"})
        except Exception as e:
            fail_count += 1
            log_entries.append({"path": raw_path, "status": "ERROR", "message": str(e)})

    return jsonify({
        "success": True,
        "total_attempted": len(paths_to_harden),
        "success_count": success_count,
        "fail_count": fail_count,
        "denied_count": denied_count,
        "backup_created": backup_path is not None,
        "logs": log_entries
    })

# -------------------------------------------------------------
# Backup Registry API
# -------------------------------------------------------------
def create_backup_file(paths, output_file):
    unique_keys = set()
    if paths:
        for p in paths:
            try:
                location, full_key = p.split("\\", 1)
                key_path, _ = full_key.split(":", 1)
                hive_full = "HKEY_LOCAL_MACHINE" if location == "HKLM" else "HKEY_CURRENT_USER" if location == "HKCU" else location
                unique_keys.add(f"{hive_full}\\{key_path}")
            except Exception:
                pass
    else:
        unique_keys = {
            "HKEY_LOCAL_MACHINE\\SYSTEM\\CurrentControlSet\\Control",
            "HKEY_LOCAL_MACHINE\\SOFTWARE\\Policies"
        }

    with open(output_file, "w", encoding="utf-16") as final_f:
        final_f.write("Windows Registry Editor Version 5.00\n\n")

    temp_dir = tempfile.mkdtemp()
    try:
        for idx, key in enumerate(unique_keys):
            temp_file = os.path.join(temp_dir, f"temp_{idx}.reg")
            cmd = f'reg.exe export "{key}" "{temp_file}" /y'
            subprocess.run(cmd, shell=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            if os.path.exists(temp_file):
                try:
                    with open(temp_file, "r", encoding="utf-16") as temp_f:
                        content = temp_f.read().replace("Windows Registry Editor Version 5.00", "").strip()
                        if content:
                            with open(output_file, "a", encoding="utf-16") as final_f:
                                final_f.write(content + "\n\n")
                except Exception:
                    pass
    finally:
        import shutil
        shutil.rmtree(temp_dir, ignore_errors=True)

@app.route("/api/backup-registry", methods=["GET", "POST"])
def backup_registry():
    backup_file = os.path.join(OUTPUTS_DIR, "backup_registry.reg")
    paths = []

    if request.method == "POST" and request.is_json:
        data = request.get_json() or {}
        paths = data.get("paths", [])
    elif STATE.get("audit_records"):
        paths = [r.get("path") for r in STATE["audit_records"] if r.get("path")]

    try:
        create_backup_file(paths, backup_file)
    except Exception as e:
        if request.method == "POST" and request.is_json:
            return jsonify({"success": False, "error": str(e)}), 500

    # If accessed directly via browser URL (GET)
    if request.method == "GET" or request.args.get("download") == "true":
        if os.path.exists(backup_file):
            return send_file(backup_file, as_attachment=True, download_name="backup_registry.reg")
        return jsonify({"error": "Failed to create registry backup file."}), 500

    # If called via POST AJAX
    return jsonify({
        "success": True,
        "path": backup_file,
        "filename": "backup_registry.reg",
        "message": "Registry backup created successfully!"
    })

# -------------------------------------------------------------
# Analytics & Chart Data API
# -------------------------------------------------------------
@app.route("/api/analytics", methods=["GET"])
def get_analytics():
    records = STATE.get("audit_records")
    csv_path = os.path.join(OUTPUTS_DIR, "registry_check_results.csv")

    if not records and os.path.exists(csv_path):
        dfr1, dfr2, dfr3 = dh.process_data(csv_path)
    else:
        dfr1, dfr2, dfr3 = dh.process_data(csv_path)

    # 1. Detection rate / Pass vs Fail
    total = len(dfr1)
    pass_count = len(dfr1[dfr1["found"].astype(str).str.lower().isin(["yes", "pass", "compliant"])])
    fail_count = total - pass_count

    # 2. Categories breakdown
    cat_counts = dfr1["Category"].value_counts().to_dict() if "Category" in dfr1 else {}
    
    # 3. System categories (dfr2)
    sys_subcats = {}
    if not dfr2.empty and "subCategory" in dfr2:
        for sub, grp in dfr2.groupby("subCategory"):
            p = len(grp[grp["found"].astype(str).str.lower().isin(["yes", "pass"])])
            f = len(grp) - p
            sys_subcats[sub] = {"pass": p, "fail": f, "total": len(grp)}

    # 4. Software categories (dfr3)
    sw_subcats = {}
    if not dfr3.empty and "subCategory" in dfr3:
        for sub, grp in dfr3.groupby("subCategory"):
            p = len(grp[grp["found"].astype(str).str.lower().isin(["yes", "pass"])])
            f = len(grp) - p
            sw_subcats[sub] = {"pass": p, "fail": f, "total": len(grp)}

    return jsonify({
        "success": True,
        "detection_rate": {
            "pass": pass_count,
            "fail": fail_count,
            "total": total,
            "score": round((pass_count / total * 100.0), 1) if total > 0 else 0.0
        },
        "categories": cat_counts,
        "system_subcategories": sys_subcats,
        "software_subcategories": sw_subcats
    })

# -------------------------------------------------------------
# File Download Endpoints (Generates & Delivers on Click)
# -------------------------------------------------------------
@app.route("/api/export-excel", methods=["GET"])
def export_excel_download():
    excel_path = os.path.join(OUTPUTS_DIR, "CIS_Benchmark_Audit_Results.xlsx")
    
    records = STATE.get("audit_records")
    if not records:
        try:
            run_audit()
            records = STATE.get("audit_records")
        except Exception as e:
            print(f"Auto-audit generation error: {e}")

    if records:
        compliance_target = STATE.get("compliance_target") or STATE.get("benchmark_title") or "CIS Security Benchmark"
        evaluation_engine = STATE.get("evaluation_engine") or "Risk Nexus CIS Security Benchmark Engine"
        metadata = {
            "hostname": platform.node() or "Local Machine",
            "os_name": f"{platform.system()} {platform.release()} ({platform.version()})",
            "scope": STATE.get("active_hive", "HKEY_LOCAL_MACHINE"),
            "compliance_target": compliance_target,
            "benchmark_title": compliance_target,
            "evaluation_engine": evaluation_engine,
            "scan_date": STATE.get("audit_timestamp") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        export_audit_results_to_excel(records, metadata, excel_path)

    if os.path.exists(excel_path):
        return send_file(excel_path, as_attachment=True, download_name="CIS_Benchmark_Audit_Results.xlsx")
    return jsonify({"error": "Audit report not found. Please run a security audit first."}), 404

@app.route("/api/export-rules-excel", methods=["GET"])
def export_rules_excel_download():
    excel_path = os.path.join(OUTPUTS_DIR, "CIS_Windows_11_Hardening_Rules.xlsx")
    if not os.path.exists(excel_path):
        if not STATE.get("rules"):
            try:
                parse_benchmark()
            except Exception as e:
                print(f"Auto-rules generation error: {e}")
        elif STATE.get("rules"):
            export_benchmark_rules_to_excel(STATE["rules"], excel_path)

    if os.path.exists(excel_path):
        return send_file(excel_path, as_attachment=True, download_name="CIS_Windows_11_Hardening_Rules.xlsx")
    return jsonify({"error": "Rules file not found. Please parse benchmark rules first."}), 404

@app.route("/api/export-csv", methods=["GET"])
def export_csv_download():
    csv_path = os.path.join(OUTPUTS_DIR, "registry_check_results.csv")
    if os.path.exists(csv_path):
        return send_file(csv_path, as_attachment=True, download_name="registry_check_results.csv")
    return jsonify({"error": "CSV file not found. Run audit first."}), 404

@app.route("/api/export-backup", methods=["GET"])
def export_backup_download():
    backup_path = os.path.join(OUTPUTS_DIR, "backup_registry.reg")
    if not os.path.exists(backup_path):
        create_backup_file([], backup_path)

    if os.path.exists(backup_path):
        return send_file(backup_path, as_attachment=True, download_name="backup_registry.reg")
    return jsonify({"error": "Backup file not found. Take backup first."}), 404

def get_local_ip():
    import socket
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"

def start_server(port=5000, host="0.0.0.0", debug=False, open_browser=True):
    import webbrowser
    import threading

    local_ip = get_local_ip()
    
    print("\n" + "=" * 72)
    print("                      RISK NEXUS SECURITY SUITE")
    print("             CIS Benchmark Audit & Remediation Platform")
    print("                 Created By Dibyataru Chakraborty")
    print("=" * 72)
    print(f" [*] Localhost Dashboard:  http://127.0.0.1:{port}")
    print(f" [*] Local Network Access: http://{local_ip}:{port}")
    print("=" * 72 + "\n")

    if open_browser:
        def _open():
            import time
            time.sleep(1.2)
            try:
                webbrowser.open(f"http://127.0.0.1:{port}")
            except Exception:
                pass
        threading.Thread(target=_open, daemon=True).start()

    app.run(host=host, port=port, debug=debug, use_reloader=False, threaded=True)

